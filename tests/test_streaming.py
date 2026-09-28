import json
import os
from pathlib import Path
import tempfile
import time
import unittest
from unittest.mock import patch
from openpyxl import load_workbook
import test_tool
from paper_citation_pipeline.streaming import StreamingExport
from paper_citation_pipeline.engine.runtime.docling_worker import Worker


def fake_worker(conn, ocr):
    while True:
        request=conn.recv()
        if request is None:break
        pdf, output, log = request
        if pdf == 'timeout':time.sleep(10)
        if pdf == 'crash':os._exit(1)
        Path(output).write_text(json.dumps(dict(status='success',page_count=os.getpid(),errors=[])))
        conn.send({'ok':True})


class StreamingTests(unittest.TestCase):
    setUp=test_tool.ToolTests.setUp
    tearDown=test_tool.ToolTests.tearDown
    write=test_tool.ToolTests.write
    packet=test_tool.ToolTests.packet
    template=test_tool.ToolTests.template

    def test_first_paper_durable_before_second_and_save_failure_preserves_it(self):
        template,mapping=self.template();original=template.read_bytes();out=self.root/'filled.xlsx'
        c=dict(template=str(template),mapping=str(mapping),output=str(out),id_state=str(self.state))
        stream=StreamingExport(c,3)
        one=self.write('one.json',self.packet());stream.advance(one)
        self.assertEqual(load_workbook(out).active.max_row,2)
        progress=json.loads(out.with_suffix('.progress.json').read_text())
        self.assertEqual((progress['processed'],progress['total']),(1,3))
        before=out.read_bytes()
        d=self.packet();d['paper_id']='P2';d['source']['sha256']='second';d['source']['pdf']='second.pdf'
        two=self.write('two.json',d)
        with patch('openpyxl.workbook.workbook.Workbook.save',side_effect=OSError('locked')):
            with self.assertRaises(OSError):stream.advance(two)
        self.assertEqual(out.read_bytes(),before)
        d=self.packet();d['paper_id']='P3';d['source']['sha256']='third';d['source']['pdf']='third.pdf'
        stream.advance(self.write('three.json',d));stream.finish()
        self.assertEqual(load_workbook(out).active.max_row,4)
        self.assertEqual(template.read_bytes(),original)

    def test_failed_parser_not_exported(self):
        template,mapping=self.template();out=self.root/'filled.xlsx'
        stream=StreamingExport(dict(template=str(template),mapping=str(mapping),output=str(out),id_state=str(self.state)),2)
        stream.advance(self.root/'missing.json',True)
        self.assertFalse(out.exists())
        stream.advance(self.write('good.json',self.packet()))
        self.assertEqual(load_workbook(out).active.max_row,2)


class WorkerTests(unittest.TestCase):
    def test_reuses_process_then_recovers_after_timeout_and_crash(self):
        with tempfile.TemporaryDirectory() as folder:
            out=Path(folder)/'doc.json';log=Path(folder)/'log'
            worker=Worker(target=fake_worker)
            try:
                a=worker.run('one',out,log,5)
                b=worker.run('two',out,log,5)
                self.assertEqual(a['page_count'],b['page_count'])
                with self.assertRaises(TimeoutError):worker.run('timeout',out,log,0.05)
                c=worker.run('three',out,log,5)
                self.assertNotEqual(a['page_count'],c['page_count'])
                with self.assertRaises((EOFError,ConnectionError)):worker.run('crash',out,log,5)
                self.assertEqual(worker.run('four',out,log,5)['status'],'success')
            finally:worker.close()

class SnapshotReuseTests(unittest.TestCase):
    setUp=test_tool.ToolTests.setUp
    tearDown=test_tool.ToolTests.tearDown
    write=test_tool.ToolTests.write
    packet=test_tool.ToolTests.packet
    template=test_tool.ToolTests.template
    def test_snapshot_reuse_requires_complete_unchanged_inputs(self):
        from paper_citation_pipeline.streaming import reuse_completed_snapshot
        template,mapping=self.template();out=self.root/'filled.xlsx'
        config=dict(template=str(template),mapping=str(mapping),output=str(out),id_state=str(self.state))
        one=self.write('one.json',self.packet());stream=StreamingExport(config,1);stream.advance(one)
        reuse=lambda:reuse_completed_snapshot([one],template,mapping,out,self.state)
        self.assertIsNone(reuse());stream.finish()
        self.assertTrue(reuse()['reused_completed_stream_snapshot'])
        for path in [one,template,mapping,self.state,out,out.with_suffix('.provenance.json'),out.with_suffix('.fill-manifest.json')]:
            original=path.read_bytes();path.write_bytes(original+b' ')
            self.assertIsNone(reuse(),str(path));path.write_bytes(original)
        self.assertIsNone(reuse_completed_snapshot([],template,mapping,out,self.state))
        self.assertTrue(reuse()['reused_completed_stream_snapshot'])

    def test_partial_parse_failure_reuses_only_saved_papers(self):
        from paper_citation_pipeline.streaming import reuse_completed_snapshot
        template,mapping=self.template();out=self.root/'filled.xlsx'
        one=self.write('one.json',self.packet())
        stream=StreamingExport(dict(template=str(template),mapping=str(mapping),output=str(out),id_state=str(self.state)),2)
        stream.advance(one);stream.advance(self.root/'missing.json',True);stream.finish()
        self.assertIsNotNone(reuse_completed_snapshot([one],template,mapping,out,self.state))

    def test_language_exclusion_is_not_exported_or_counted_as_failure(self):
        template,mapping=self.template();out=self.root/'filled.xlsx'
        stream=StreamingExport(dict(template=str(template),mapping=str(mapping),output=str(out),id_state=str(self.state)),2)
        before=self.state.read_bytes()
        stream.advance(self.root/'portuguese.json',skipped=True)
        self.assertFalse(out.exists());self.assertEqual(before,self.state.read_bytes())
        stream.advance(self.write('english.json',self.packet()));stream.finish()
        progress=json.loads(out.with_suffix('.progress.json').read_text())
        self.assertEqual((progress['processed'],progress['succeeded'],progress['excluded'],progress['failed']),(2,1,1,0))
        events=json.loads(out.with_suffix('.checkpoints.json').read_text())['events']
        self.assertEqual(events[0]['status'],'excluded')

    def test_all_non_english_batch_exports_empty_template_and_exclusion_log(self):
        import pymupdf
        from test_language import PT
        from paper_citation_pipeline.cli import main
        template,mapping=self.template();cache=self.root/'cache';(cache/'json').mkdir(parents=True)
        pdf=self.root/'portuguese.pdf';doc=pymupdf.open()
        for _ in range(3):doc.new_page().insert_textbox(pymupdf.Rect(40,100,550,700),PT,fontsize=11)
        doc.save(pdf);doc.close()
        (cache/'json/portuguese.json').write_text(json.dumps({'source':{'pdf':str(pdf)}}))
        before=self.state.read_bytes();out=self.root/'results'
        status=main(['run','--parser-output',str(cache),'--output',str(out),'--target','UNEP',
            '--id-state',str(self.state),'--template',str(template),'--mapping',str(mapping)])
        self.assertEqual(status,0);self.assertEqual(before,self.state.read_bytes())
        self.assertFalse(list((out/'json').glob('*.json')))
        self.assertEqual(json.loads((out/'manifest.json').read_text())[0]['readiness'],'excluded')
        self.assertEqual(json.loads((out/'citations.provenance.json').read_text())['counts']['citations'],0)
        issues=json.loads((out/'exceptions.json').read_text())['issues']
        self.assertEqual(issues[0]['kind'],'非英语论文已排除')
        self.assertFalse(any(i['severity']=='error' for i in issues))

        progress=json.loads((out/'citations.progress.json').read_text())
        self.assertEqual((progress['succeeded'],progress['excluded'],progress['failed']),(0,1,0))
