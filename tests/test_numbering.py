import contextlib
import io
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
import test_tool
from paper_citation_pipeline.numbering import initialize, assign, check, locked
from paper_citation_pipeline.excel import export_packets
from paper_citation_pipeline.cli import main
from paper_citation_pipeline.progress import Progress


class NumberingTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.state=self.root/'numbering.json';initialize(self.state,25,100,'P','R')
    def tearDown(self):self.tmp.cleanup()
    def packet(self):return test_tool.ToolTests.packet(self)
    def test_last_used_and_repeat(self):
        a=assign(self.packet(),self.state)
        self.assertEqual(a['numbering']['paper_id'],'P000026')
        self.assertEqual(a['numbering']['records']['e1'],'R000101')
        self.assertEqual(a['paper_id'],'P1');self.assertEqual(a['target_candidates'][0]['record_id'],'e1')
        b=assign(self.packet(),self.state)
        self.assertEqual(a['numbering'],b['numbering'])
        self.assertEqual(check(self.state)['record']['current'],101)
    def test_new_target_reuses_paper_but_not_record(self):
        a=assign(self.packet(),self.state)
        b=self.packet();b['target']['canonical_name']='UNEP';assign(b,self.state)
        self.assertEqual(a['numbering']['paper_id'],b['numbering']['paper_id'])
        self.assertEqual(b['numbering']['records']['e1'],'R000102')
    def test_zero_citations_still_get_paper_id(self):
        a=self.packet();a['target_candidates']=[];assign(a,self.state)
        self.assertEqual(check(self.state)['record']['current'],100)
        self.assertEqual(check(self.state)['paper']['current'],26)
    def test_conflicts_get_separate_record(self):
        a=self.packet();a['unresolved']={'source_consistency_conflicts':[{'location_id':'l2'}]};assign(a,self.state)
        self.assertEqual(a['numbering']['records']['conflict_l2'],'R000102')
    def test_reset_duplicate_and_other_ledger_rejected(self):
        with self.assertRaises(FileExistsError):initialize(self.state,0,0)
        a=assign(self.packet(),self.state);other=self.root/'other.json';initialize(other,0,0)
        before=other.read_bytes()
        with self.assertRaises(ValueError):assign(a,other)
        self.assertEqual(before,other.read_bytes())
        state=check(self.state);state['record']['current']=0;self.state.write_text(json.dumps(state))
        with self.assertRaises(ValueError):check(self.state)
    def test_lock_blocks_other_writer(self):
        with locked(self.state):
            with self.assertRaises(ValueError):assign(self.packet(),self.state)
    def test_missing_state_blocks_engine(self):
        with patch('paper_citation_pipeline.cli.subprocess.run') as run:
            self.assertEqual(main(['run','--target','UNEP']),2)
        run.assert_not_called()
    def test_failed_save_reserves_and_retry_reuses(self):
        from openpyxl import Workbook,load_workbook
        wb=Workbook();wb.active.title='data';wb.active.append(['Paper','Record'])
        template=self.root/'blank.xlsx';wb.save(template)
        mapping=self.root/'map.json';mapping.write_text(json.dumps({'sheets':[{'sheet':'data','entity':'citations','columns':{'Paper':'paper_id','Record':'record_id'}}]}))
        packet=self.root/'packet.json';packet.write_text(json.dumps(self.packet()))
        output=self.root/'filled.xlsx';log=io.StringIO()
        with contextlib.redirect_stderr(log),patch('openpyxl.workbook.workbook.Workbook.save',side_effect=OSError('disk full')):
            with self.assertRaises(OSError):export_packets([packet],template,mapping,output,id_state=self.state)
        self.assertFalse(output.exists());self.assertNotIn('1/1',log.getvalue());self.assertIn('failed',log.getvalue())
        result=export_packets([packet],template,mapping,output,id_state=self.state)
        self.assertEqual(result['record_ids'],['R000101'])
        self.assertEqual(check(self.state)['record']['current'],101)
        self.assertEqual(load_workbook(output).active['A2'].value,'P000026')
    def test_progress_counts_failures_and_persists(self):
        p=self.root/'progress.json';bar=Progress('PDF',2,p);bar.advance('one');bar.advance('two',True);bar.finish()
        d=json.loads(p.read_text());self.assertEqual((d['processed'],d['total'],d['succeeded'],d['failed']),(2,2,1,1))
        self.assertEqual(d['status'],'completed_with_failures')

if __name__=='__main__':unittest.main()
