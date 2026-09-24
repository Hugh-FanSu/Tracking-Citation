import json,sys,tempfile,unittest
from pathlib import Path
from unittest.mock import patch,Mock
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from paper_citation_pipeline.ai_review import config,Client
from paper_citation_pipeline.quality import inspect_workbook,review_workbook
from paper_citation_pipeline.excel import export_packets
from paper_citation_pipeline.numbering import initialize
from openpyxl import load_workbook
import test_tool

class AIReviewTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.c=self.root/'ai.json';self.c.write_text('{}');self.cfg=config(self.c)
        self.state=self.root/'state.json';initialize(self.state,0,0)
        self.packet=test_tool.ToolTests.packet(self)
        self.packet['metadata']['abstract']='PRIVATE_ABSTRACT_CANARY'
        self.packet['target_candidates'][0]['paragraph_text']='PRIVATE_PARAGRAPH_CANARY'
        self.path=self.root/'paper.json';self.path.write_text(json.dumps(self.packet))
        self.template,self.mapping=test_tool.ToolTests.template(self)
        self.out=self.root/'filled.xlsx'
        export_packets([self.path],self.template,self.mapping,self.out,id_state=self.state)
    def write(self,name,data):
        p=self.root/name;p.write_text(json.dumps(data));return p
    def tearDown(self):self.temp.cleanup()
    def test_saved_workbook_evidence_detects_tampered_cell(self):
        from paper_citation_pipeline.stage_monitor import filled_summary
        before=filled_summary(self.out)
        self.assertEqual(before['cells_compared'],before['cells_matched'])
        wb=load_workbook(self.out);wb.active['B2']='corrupted-value';wb.save(self.out);wb.close()
        after=filled_summary(self.out)
        self.assertGreater(after['issue_counts']['cell_mismatch'],0)
        self.assertLess(after['cells_matched'],after['cells_compared'])
        self.assertNotEqual(before['workbook_sha256'],after['workbook_sha256'])
    def test_config_rejects_keys_and_remote_http(self):
        for data in [{'api_key':'secret'},{'endpoint':'http://example.org/v1'}]:
            self.c.write_text(json.dumps(data))
            with self.assertRaises(ValueError):config(self.c)
    def test_valid_export_and_source_not_needed(self):
        self.path.unlink()
        self.assertEqual(review_workbook(self.out),0)
        report=json.loads(self.out.with_suffix('.quality.json').read_text())
        self.assertEqual(report['status'],'passed')
    def test_lost_cell_and_dropped_row_detected(self):
        wb=load_workbook(self.out);wb.active['B2']=None;wb.save(self.out)
        _,checks=inspect_workbook(self.out)
        self.assertTrue(any(c['kind']=='cell_mismatch' for c in checks))
        wb.active.delete_rows(2);wb.save(self.out)
        _,checks=inspect_workbook(self.out)
        self.assertTrue(any(c['kind']=='row_count' and c['status']=='error' for c in checks))
    def test_extra_row_and_broken_link(self):
        wb=load_workbook(self.out);wb.active.append(['UNKNOWN','value','待核实']);wb.save(self.out)
        _,checks=inspect_workbook(self.out)
        self.assertTrue(any(c['kind']=='broken_foreign_key' for c in checks))
        self.assertEqual(review_workbook(self.out),1)
    def test_input_paper_omission(self):
        self.assertEqual(review_workbook(self.out,expected_papers=2),1)
    def test_legacy_ai_config_cannot_trigger_model(self):
        with patch('paper_citation_pipeline.ai_review.Client') as client:
            self.assertEqual(review_workbook(self.out,self.cfg),0)
            client.assert_not_called()
        self.assertFalse(json.loads(self.out.with_suffix('.quality.json').read_text())['ai_enabled'])
    def test_duplicate_ids_and_incorrect_mapping(self):
        from openpyxl import Workbook
        wb=Workbook();wb.active.title='引用记录';wb.active.append(['记录编号*','论文编号*','所在完整段落*'])
        wb.save(self.template)
        self.mapping.write_text(json.dumps({'sheets':[{'sheet':'引用记录','entity':'citations','columns':{'记录编号*':'record_id','论文编号*':'paper_id','所在完整段落*':'raw_marker'}}]}))
        export_packets([self.path],self.template,self.mapping,self.out,True,id_state=self.state)
        wb=load_workbook(self.out);wb.active.append([wb.active['A2'].value,wb.active['B2'].value,'value']);wb.save(self.out)
        _,checks=inspect_workbook(self.out)
        self.assertTrue(any(c['kind']=='duplicate_identifiers' and c['status']=='error' for c in checks))
        self.assertTrue(any(c['kind']=='field_mapping' and c['status']=='error' for c in checks))
    def test_missing_required_source_field_distinguished(self):
        from openpyxl import Workbook
        wb=Workbook();wb.active.title='引用记录';wb.active.append(['论文编号*','所在完整段落*']);wb.save(self.template)
        self.mapping.write_text(json.dumps({'sheets':[{'sheet':'引用记录','entity':'citations','columns':{'论文编号*':'paper_id','所在完整段落*':'paragraph_text'}}]}))
        self.packet['target_candidates'][0]['paragraph_text']=None
        self.path.write_text(json.dumps(self.packet))
        export_packets([self.path],self.template,self.mapping,self.out,True,id_state=self.state)
        _,checks=inspect_workbook(self.out)
        self.assertTrue(any(c['kind']=='required_blank_in_source' for c in checks))
        self.assertFalse(any(c['kind']=='cell_mismatch' for c in checks))
    def test_legacy_semantic_ai_is_not_exported(self):
        from paper_citation_pipeline.excel import collect
        self.packet['ai_review']={'status':'completed','findings':[], 'decisions':[{'record_id':'e1','verdict':'supports','reason':'OLD_AI_CANARY','evidence':[]}]}
        self.path.write_text(json.dumps(self.packet))
        data=collect([self.path],self.state)
        self.assertNotIn('OLD_AI_CANARY',json.dumps(data,default=str))
        self.assertEqual(data['citations'][0]['confirmation_status'],'待核实')
    def test_client_auth_request_json_and_rejects_redirects(self):
        client=Client(self.cfg,'secret');resp=Mock(status_code=200)
        resp.json.return_value={'choices':[{'finish_reason':'stop','message':{'content':'{"ok":true}'}}]}
        client.session.post=Mock(return_value=resp);self.assertIn('成功',client.test())
        kwargs=client.session.post.call_args.kwargs
        self.assertEqual(kwargs['json']['thinking'],{'type':'disabled'})
        self.assertEqual(kwargs['headers']['Authorization'],'Bearer secret');self.assertFalse(kwargs['allow_redirects'])
        self.assertEqual(kwargs['json']['response_format'],{'type':'json_object'})
        resp.status_code=302
        with self.assertRaises(ValueError):client.test()
        client.close()
    def test_truncated_model_response_is_not_success(self):
        client=Client(self.cfg,'secret');resp=Mock(status_code=200)
        resp.json.return_value={'usage':{'total_tokens':1536},'choices':[{'finish_reason':'length','message':{'content':'{"ok":true}'}}]}
        client.session.post=Mock(return_value=resp)
        with self.assertRaises(ValueError):client.test()
        self.assertEqual(client.last_metadata['usage']['total_tokens'],1536)
        client.close()
    def test_explicit_diagnostic_is_single_small_request(self):
        client=Client(self.cfg,'secret');client.session.post=Mock(return_value=Mock(status_code=429))
        with self.assertRaises(ValueError):client.test()
        self.assertEqual(client.session.post.call_count,1)
        self.assertEqual(client.session.post.call_args.kwargs['json']['max_tokens'],64)
        client.close()
    def test_retry_rate_limits_is_bounded(self):
        client=Client(self.cfg,'secret');client.session.post=Mock(return_value=Mock(status_code=429))
        with patch('paper_citation_pipeline.ai_review.time.sleep'):
            with self.assertRaises(ValueError):client.request([])
        self.assertEqual(client.session.post.call_count,3);client.close()

if __name__=="__main__":unittest.main()
