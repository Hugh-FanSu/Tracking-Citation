import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch
from openpyxl import Workbook
from paper_citation_pipeline.cli import main
from paper_citation_pipeline.config import read_json
from paper_citation_pipeline.numbering import initialize
from paper_citation_pipeline.setup import template_mapping


class SetupTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory();self.root=Path(self.tmp.name)
        self.pdfs=self.root/'PDFs';self.pdfs.mkdir();(self.pdfs/'sample.PDF').write_bytes(b'%PDF-fixture')
        self.template=self.root/'blank.xlsx';w=Workbook();w.active.title='citations';w.active.append(['record_id','paper_id','paragraph_text']);w.save(self.template)
        self.project=self.root/'project'
        self.flags=['setup',str(self.project),'--input',str(self.pdfs),'--template',str(self.template),'--target','WHO','--paper-current','20','--record-current','80','--non-interactive']
    def tearDown(self):self.tmp.cleanup()
    def test_new_project_is_portable_and_originals_unchanged(self):
        before=self.template.read_bytes();self.assertEqual(main(self.flags),0)
        cfg=read_json((self.project/'run.json').resolve())
        self.assertEqual(cfg['input'],'../PDFs');self.assertEqual(cfg['template'],'assets/template.xlsx')
        self.assertEqual(read_json(self.project/'numbering.json')['record']['current'],80)
        self.assertEqual(before,self.template.read_bytes());self.assertEqual(before,(self.project/cfg['template']).read_bytes())
        self.assertTrue((self.project/'compose.yaml').is_file())
        self.assertEqual(main(self.flags),2)
    def test_missing_counters_do_not_default_to_zero(self):
        flags=self.flags[:self.flags.index('--paper-current')]+['--non-interactive']
        self.assertEqual(main(flags),2);self.assertFalse(self.project.exists())
    def test_unknown_column_requires_explicit_mapping(self):
        w=Workbook();w.active.title='citations';w.active.append(['mystery']);w.save(self.template)
        self.assertEqual(main(self.flags),2);self.assertFalse(self.project.exists())
        with patch('builtins.input',side_effect=['1','2','paragraph_text']):
            spec=template_mapping(self.template,True)
        self.assertEqual(spec['sheets'][0]['columns'],{'mystery':'paragraph_text'})
    def test_custom_template_explicit_mapping(self):
        w=Workbook();w.active.title='MySheet';w.active.append(['正文']);w.save(self.template)
        m=self.root/'mapping.json';m.write_text(json.dumps({'sheets':[{'sheet':'MySheet','entity':'citations','columns':{'正文':'paragraph_text'}}]}))
        self.assertEqual(main(self.flags+['--mapping',str(m)]),0)
    def test_existing_state_is_referenced_not_forked(self):
        state=self.root/'ledger.json';initialize(state,50,500)
        flags=self.flags[:self.flags.index('--paper-current')]+['--non-interactive','--id-state',str(state)]
        self.assertEqual(main(flags),0)
        self.assertEqual(read_json((self.project/'run.json').resolve())['id_state'],'../ledger.json')
        self.assertFalse((self.project/'numbering.json').exists())
    def test_invalid_mapping_does_not_leave_partial_project(self):
        m=self.root/'map.json';m.write_text(json.dumps({'sheets':[{'sheet':'absent','entity':'citations','columns':{'id':'record_id'}}]}))
        self.assertEqual(main(self.flags+['--mapping',str(m)]),2)
        self.assertFalse(self.project.exists());self.assertFalse(list(self.root.glob('.paper-citations-setup-*')))
    def test_setup_run_uses_public_pipeline(self):
        with patch('paper_citation_pipeline.cli.run',return_value=1) as run:
            self.assertEqual(main(self.flags+['--run']),1)
        self.assertEqual(Path(run.call_args.args[0].config),(self.project/'run.json').resolve())

if __name__=='__main__':unittest.main()
