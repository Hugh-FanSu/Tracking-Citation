import hashlib
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from types import SimpleNamespace
from unittest.mock import patch
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'src'))
from paper_citation_pipeline.config import target_config
from paper_citation_pipeline.excel import export_packets,load_template,visual_context
from paper_citation_pipeline.cli import main
from paper_citation_pipeline.numbering import initialize
from openpyxl import Workbook,load_workbook
ROOT=Path(__file__).resolve().parents[1]

class ToolTests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.root=Path(self.temp.name)
        self.state=self.root/"numbering.json"
        initialize(self.state,10,100)
    def tearDown(self):self.temp.cleanup()
    def write(self,name,data):
        path=self.root/name;path.write_text(json.dumps(data));return path
    def packet(self):
        return {'paper_id':'P1','source':{'pdf':'original.pdf','sha256':'source'},'page_count':1,
                'target':{'canonical_name':'WHO'},'quality':{'readiness':'ready_with_issues'},'metadata':{'title':'Study','authors':[],'abstract':'Abstract'},
                'parser_result':{'software':{}},'citation_index':{'references':[{'reference_id':'r1','raw_citation':'WHO, 2020. Report.'}],
                'contexts':[{'id':'p1','text':'=not_a_formula','kind':'p'}],'sections':[],
                'occurrences':[{'location_id':'loc1','links':[{'reference_id':'r1','local_reference_id':'b1'}]}]},
                'target_candidates':[{'record_id':'e1','paper_id':'P1','location_id':'loc1','paragraph_id':'p1','paragraph_text':'=not_a_formula','raw_marker':'(WHO, 2020)',
                 'raw_reference':'WHO, 2020. Report.','report_candidate_id':'r1','report_match':{'status':'not_searched'},'link_status':'machine_linked'}]}
    def template(self):
        wb=Workbook();ws=wb.active;ws.title='Citations';ws.append(['Paper','Text','Status']);ws.freeze_panes='B2'
        ws['A2'].number_format='@';path=self.root/'blank.xlsx';wb.save(path)
        mapping=self.write('map.json',{'sheets':[{'sheet':'Citations','entity':'citations','columns':{'Paper':'paper_id','Text':'paragraph_text','Status':'confirmation_status'}}]})
        return path,mapping
    def test_alias_metadata_must_match_target(self):
        with self.assertRaises(ValueError):target_config('WHO',ROOT/'examples/unep-aliases.json')
    def test_literal_list_and_kind(self):
        p=self.write('alias.json',['W.H.O.','World Health Organization'])
        d=target_config('WHO',p);self.assertEqual(d['aliases'][0]['name'],'WHO')
        with self.assertRaises(ValueError):target_config('WHO',kind='unknown')
    def test_targets_isolated_in_subprocesses(self):
        code="""import json,sys
sys.path.insert(0,sys.argv[1])
from standardize_citations import normalize,unep_rows
d={'source':{'sha256':'x','pdf':'x'},'sections':[],'citation_mentions':[],
'paragraphs':[{'id':'p','text':'(WHO, 2020); (UNEP, 2020); (Example et al., 2021)','kind':'p','coordinates':[]}],
'bibliography':[{'id':'a','raw_citation':'WHO, 2020. Report.','year':'2020'},
{'id':'b','raw_citation':'UNEP, 2020. Report. https://unep.org/report','year':'2020'},
{'id':'c','raw_citation':'Example, J., 2021. Study.','year':'2021'}]}
print(json.dumps([r['organization_id'] for r in unep_rows(normalize(d,'p'))]))
"""
        for name,file,expected in [('WHO','who-aliases.json','WHO'),('Jane Example','person-aliases.json','jane-example'),('UNEP','unep-aliases.json','UNEP')]:
            target=self.write('target.json',target_config(name,ROOT/'examples'/file))
            run=subprocess.run([sys.executable,'-c',code,str(ROOT/'src/paper_citation_pipeline/engine/runtime')],env=dict(os.environ,PAPER_CITATION_TARGET=str(target)),capture_output=True,text=True,check=True)
            self.assertEqual(json.loads(run.stdout),[expected])
    def test_export_keeps_template_and_literal_pdf_text(self):
        template,mapping=self.template();before=template.read_bytes();packet=self.write('p.json',self.packet());out=self.root/'filled.xlsx'
        result=export_packets([packet],template,mapping,out,id_state=self.state)
        self.assertEqual(result['counts']['citations'],1);self.assertEqual(before,template.read_bytes())
        ws=load_workbook(out).active
        self.assertEqual(ws['B2'].value,'=not_a_formula');self.assertEqual(ws['B2'].data_type,'s');self.assertEqual(ws['C2'].value,'待核实');self.assertEqual(ws.freeze_panes,'B2')
        with self.assertRaises(FileExistsError):export_packets([packet],template,mapping,out,id_state=self.state)
    def test_cannot_overwrite_template(self):
        t,m=self.template()
        with self.assertRaises(ValueError):export_packets([],t,m,t,True,id_state=self.state)
    def test_mapping_and_occupied_cells_fail(self):
        t,m=self.template();spec=json.loads(m.read_text());spec['sheets'][0]['columns']['Missing']='paper_id';m.write_text(json.dumps(spec))
        with self.assertRaises(ValueError):load_template(t,m)
        t,m=self.template();w=load_workbook(t);w.active['A2']='Existing';w.save(t)
        with self.assertRaises(ValueError):export_packets([self.write('p.json',self.packet())],t,m,self.root/'filled.xlsx',id_state=self.state)
    def test_long_text_never_silently_truncated(self):
        t,m=self.template();d=self.packet();d['target_candidates'][0]['paragraph_text']='x'*32768
        with self.assertRaises(ValueError):export_packets([self.write('p.json',d)],t,m,self.root/'filled.xlsx',id_state=self.state)
    def test_visual_context_keeps_caption_cells_and_notes(self):
        d=self.packet();d['parser_result']['docling']={'document':{'texts':[{'self_ref':'#/texts/1','text':'Table 1: evidence'},{'self_ref':'#/texts/2','text':'Source: WHO'}],
        'tables':[{'self_ref':'#/tables/1','label':'table','captions':[{'$ref':'#/texts/1'}],'footnotes':[{'$ref':'#/texts/2'}],'data':{'table_cells':[{'text':'cell text [4]'}]}}]}}
        result=visual_context(d,{'docling_table_ids':['#/tables/1']})
        self.assertEqual(result[0]['footnotes'],['Source: WHO']);self.assertEqual(result[0]['cells'],['cell text [4]'])
    def test_invalid_template_stops_before_engine(self):
        cache=self.root/'cache';cache.mkdir();t,m=self.template();spec=json.loads(m.read_text());spec['sheets'][0]['sheet']='missing';m.write_text(json.dumps(spec))
        with patch('paper_citation_pipeline.cli.subprocess.run') as run:
            code=main(['run','--id-state',str(self.state),'--parser-output',str(cache),'--output',str(self.root/'out'),'--target','WHO','--template',str(t),'--mapping',str(m)])
        self.assertEqual(code,2);run.assert_not_called();self.assertFalse((self.root/'out/json').exists());self.assertTrue((self.root/'out/exceptions.txt').exists())
    def test_config_paths_resolve_next_to_config(self):
        cache=self.root/'cache';cache.mkdir()
        conf=self.write('run.json',{'parser_output':'cache','output':'result','target':'WHO','id_state':'numbering.json'})
        def finish(command,**kwargs):
            self.assertEqual(command[command.index('--parser-output')+1],str(cache.resolve()))
            (self.root/'result/manifest.json').write_text('[]')
            return SimpleNamespace(returncode=0)
        with patch('paper_citation_pipeline.cli.subprocess.run',side_effect=finish):
            self.assertEqual(main(['run','--config',str(conf)]),0)
        self.assertEqual(json.loads((self.root/'result/target.json').read_text())['organization_id'],'WHO')
    def test_init_and_skill_install_are_explicit(self):
        dest=self.root/'project'
        self.assertEqual(main(['init',str(dest)]),0)
        self.assertTrue((dest/'examples/unep-template-map.json').exists())
        self.assertEqual(main(['init',str(dest)]),2)
        skill=self.root/'skill'
        self.assertEqual(main(['install-skill','--path',str(skill)]),0)
        self.assertTrue((skill/'runtime.json').exists())
        self.assertEqual(main(['install-skill','--path',str(skill),'--force']),0)
        self.assertEqual(len(list(self.root.glob('skill.backup-*'))),1)
        unsafe=self.root/'not-skill';unsafe.mkdir()
        self.assertEqual(main(['install-skill','--path',str(unsafe),'--force']),2)

if __name__=='__main__':unittest.main()
