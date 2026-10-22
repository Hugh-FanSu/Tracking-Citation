import json
from pathlib import Path
import unittest
import tomllib
import openpyxl
from paper_citation_pipeline import __version__
from paper_citation_pipeline.template_bundle import load_bundle

ROOT = Path(__file__).resolve().parents[1]


class ReleaseResourcesTests(unittest.TestCase):
    def test_versions_agree(self):
        self.assertEqual(__version__, tomllib.loads((ROOT / 'pyproject.toml').read_text())['project']['version'])

    def test_template_mapping_and_blank_collection_area(self):
        bundle = load_bundle(ROOT / 'templates/UNEP', expected_org='UNEP')
        wb = openpyxl.load_workbook(bundle['template'])
        try:
            names = ['CitationRecords', 'PaperInfo', 'ReportInfo', 'EntryExamples', 'Instructions']
            self.assertEqual([(table.name, table.displayName) for sheet in wb for table in sheet.tables.values()],
                             [(name, name) for name in names])
            for sheet in list(wb)[:3]:
                self.assertFalse(any(c.value is not None for row in sheet.iter_rows(min_row=2) for c in row))
            self.assertEqual([len(s.data_validations.dataValidation) for s in list(wb)[:3]], [5, 1, 1])
            self.assertIn(__version__, wb['填写说明']['B18'].value)
        finally:
            wb.close()
        runtime = ROOT / 'src/paper_citation_pipeline/engine/runtime/config/organization_aliases.json'
        self.assertEqual(json.loads(bundle['aliases'].read_text()), json.loads(runtime.read_text()))

    def test_diagram_sources_match_markdown(self):
        for source in (ROOT / 'docs/diagrams').glob('*.mmd'):
            self.assertIn(source.read_text(), source.with_suffix('.md').read_text())
