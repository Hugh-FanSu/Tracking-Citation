"""Template-only XLSX export. No Markdown parsing or confirmation of candidates."""
from copy import copy
from datetime import datetime
import json
from pathlib import Path
import re
import tempfile
import xml.etree.ElementTree as ET
from .config import read_json, digest
from .numbering import assign, check
from .progress import Progress

FIELDS={
 'citations':set('record_id paper_id report_id location_id part section_title subsection_title pdf_pages printed_page paragraph_id occurrence_in_paragraph raw_marker target_marker citation_sentence paragraph_text previous_paragraph next_paragraph carrier raw_reference other_references confirmation_status context_status notes collector reviewer processed_date rule_version software report_match_status report_title source_json source_pdf target_name link_status visual_context'.split()),
 'papers':set('paper_id title doi eid year authors journal abstract purpose conclusion source_pdf source_sha256 file_version page_count fulltext_status batch notes source_json target_name'.split()),
 'reports':set('report_id target_name title year edition identifiers url source_pdf language verification_status notes'.split())}

def load_template(template,mapping):
    from openpyxl import load_workbook
    template=Path(template).resolve()
    if template.suffix.lower()!='.xlsx':raise ValueError('Only local .xlsx templates are supported')
    spec=read_json(mapping)
    if not isinstance(spec,dict) or not isinstance(spec.get('sheets'),list) or not spec['sheets']:
        raise ValueError('Mapping must contain a nonempty sheets list')
    wb=load_workbook(template)
    seen=set()
    for cfg in spec['sheets']:
        name=cfg['sheet'];entity=cfg['entity']
        if name not in wb.sheetnames or name in seen:raise ValueError(f'Missing or duplicate mapped sheet: {name}')
        seen.add(name)
        if entity not in FIELDS:raise ValueError(f'Unknown entity: {entity}')
        header=cfg.get('header_row',1);start=cfg.get('start_row',header+1)
        if not isinstance(header,int) or not isinstance(start,int) or header<1 or start<=header:raise ValueError('Invalid header_row/start_row')
        if not cfg.get('columns'):raise ValueError('Empty column mapping')
        headers={}
        for cell in wb[name][header]:
            if cell.value is not None:headers.setdefault(str(cell.value),[]).append(cell.column)
        for title,field in cfg['columns'].items():
            if len(headers.get(title,[]))!=1:raise ValueError(f'Header missing or duplicated: {name}/{title}')
            if field not in FIELDS[entity]:raise ValueError(f'Unknown field: {entity}.{field}')
        cfg['_positions']={headers[h][0]:f for h,f in cfg['columns'].items()}
    return wb,spec

def classify(title):
    s=(title or '').lower()
    if 'result' in s and 'discussion' in s:return '结果与讨论'
    for regex,label in [('abstract','摘要'),('intro','引言'),('review|theor|background','文献综述或理论'),('method|material|data|experimental','数据与方法'),('result','结果'),('discuss','讨论'),('conclu','结论'),('append|supplement','附录或补充材料')]:
        if re.search(regex,s):return label
    return '其他' if s else '无法判断'

def visual_context(packet,row):
    doc=packet['parser_result'].get('docling',{}).get('document',{})
    texts={x['self_ref']:x.get('text','') for x in doc.get('texts',[])}
    context=next((p for p in packet['citation_index']['contexts'] if p['id']==row.get('paragraph_id')), {})
    matches={m['docling_id'] for m in context.get('docling_matches',[])}
    matches.update(row.get('docling_table_ids',[]))
    result=[]
    for item in doc.get('tables',[])+doc.get('pictures',[]):
        related={x['$ref'] for k in ['captions','footnotes','children'] for x in item.get(k,[])}|{item['self_ref']}
        if not matches & related:continue
        result.append({'id':item['self_ref'],'label':item['label'],
                       'captions':[texts.get(x['$ref'],'') for x in item.get('captions',[])],
                       'footnotes':[texts.get(x['$ref'],'') for x in item.get('footnotes',[])],
                       'cells':[c.get('text','') for c in item.get('data',{}).get('table_cells',[])]})
    return result

def collect(paths, id_state=None):
    citations=[];papers=[];reports={};targets=set();seen=set()
    ns={'t':'http://www.tei-c.org/ns/1.0'}
    def txt(e):return ''.join(e.itertext()).strip() if e is not None else None
    for path in paths:
        d=read_json(path)
        if d.get('quality',{}).get('readiness')=='failed' or 'citation_index' not in d:raise ValueError(f'Not a usable packet: {path}')
        if id_state is not None:assign(d,id_state)
        pid=d.get('numbering',{}).get('paper_id',d['paper_id'])
        if pid in seen:raise ValueError('Duplicate paper_id: '+pid)
        seen.add(pid)
        target=d.get('target',{}).get('canonical_name') or next((r.get('organization_name_canonical') for r in d['target_candidates']),None)
        targets.add(target)
        idx=d['citation_index'];ps={p['id']:p for p in idx['contexts']};sections={s['id']:s for s in idx['sections']};refs={r['reference_id']:r for r in idx['references']};occ={o['location_id']:o for o in idx['occurrences']}
        doi=year=journal=None
        tei=d.get('provenance',{}).get('artifacts',{}).get('tei',{}).get('path')
        if tei and Path(tei).is_file():
            header=ET.parse(tei).find('.//t:teiHeader',ns)
            if header is not None:
                doi=txt(header.find('.//t:sourceDesc//t:idno[@type="DOI"]',ns));journal=txt(header.find('.//t:sourceDesc//t:monogr/t:title',ns))
                date=header.find('.//t:publicationStmt/t:date[@type="published"]',ns)
                match=re.search(r'(?:19|20)\d{2}',date.get('when','') if date is not None else '')
                year=int(match.group()) if match else None
        papers.append(dict(source_paper_id=d['paper_id'],paper_id=pid,title=d['metadata'].get('title'),doi=doi,year=year,journal=journal,
                           authors='; '.join(a.get('name','') for a in d['metadata'].get('authors') or []),abstract=d['metadata'].get('abstract'),
                           source_pdf=d['source']['pdf'],source_sha256=d['source']['sha256'],file_version='原始PDF',page_count=d['page_count'],
                           fulltext_status='检查中',batch=Path(path).parent.parent.name,source_json=str(path),target_name=target,
                           notes='结构解析已完成，尚未逐条语义确认。JSON='+str(path)))
        rows=list(d['target_candidates'])
        for conflict in d.get('unresolved',{}).get('source_consistency_conflicts',[]):
            o=occ[conflict['location_id']];p=ps[conflict['paragraph_id']];r=refs[conflict['title_suggested_reference_id']]
            rows.append(dict(record_id='conflict_'+o['location_id'],location_id=o['location_id'],paragraph_id=p['id'],raw_marker=conflict['printed_marker'],
                             paragraph_text=p['text'],previous_paragraph=ps.get(p.get('previous_paragraph_id'),{}).get('text'),next_paragraph=ps.get(p.get('next_paragraph_id'),{}).get('text'),
                             coordinates=conflict['coordinates'],carrier=p['kind'],raw_reference=r['raw_citation'],parsed_title_unverified=r.get('title'),report_candidate_id=r['reference_id'],
                             report_match={'status':'blocked_reference'},link_status='source_conflict',conflict=conflict))
        for r in rows:
            row={key:r.get(key) for key in FIELDS['citations']};row.update(source_record_id=r['record_id'],source_paper_id=d['paper_id'],paper_id=pid,source_json=str(path),source_pdf=d['source']['pdf'],target_name=target)
            if d.get('numbering'):row['record_id']=d['numbering']['records'][r['record_id']]
            sec=sections.get(ps.get(r.get('paragraph_id'),{}).get('section_id'),{});chain=[];visited=set()
            while sec and sec['id'] not in visited:
                chain.insert(0,((sec.get('number') or '')+' '+(sec.get('title') or '')).strip());visited.add(sec['id']);sec=sections.get(sec.get('parent_id'),{})
            coords=r.get('coordinates') or r.get('context_coordinates') or []
            pages=sorted({c['page'] for c in coords})
            match=r.get('report_match',{'status':'not_searched'})
            disputed=match['status']=='matched' and any('年份' in x and '不同' in x for x in match.get('notes',[]))
            rid=match.get('catalog_resource_id') if match['status']=='matched' and not disputed else None
            title=(match.get('matched_title') or 0) if match['status']=='matched' and not disputed else 0
            notes=['机器候选；未作有效内容引用确认。','link_status='+str(r.get('link_status')),'JSON='+str(path),'reference_id='+str(r.get('report_candidate_id'))]
            if r.get('author_review'):
                review=r['author_review']; result=review.get('result',{})
                notes.append('作者角色API：'+result.get('role',review.get('status','pending'))+'；'+result.get('reason',review.get('reason',review.get('error',''))))
            if r.get('joint_authorship'):
                notes.append('按共同/复合署名规则纳入；保留原始署名：'+r['joint_authorship']['raw_author_label'])
            if r.get('report_identity_status')=='candidate_related_organization_scope':
                notes.append('组织统计范围待确认：复合署名仅与目标简称共享前缀，不代表已确认是目标组织别名。原署名见原始参考文献。')
            if r.get('resource_type')=='document_identifier_unresolved':
                notes.append('表格文号来源；报告题名和出版年份未从文号推断。')
            if disputed:notes.append('日期/版本差异；原JSON名称匹配='+json.dumps(match,ensure_ascii=False))
            if r.get('conflict'):notes.append('原文数字与题名冲突，禁止自动改号：'+json.dumps(r['conflict'],ensure_ascii=False))
            if not r.get('citation_sentence'):notes.append('缺可靠引用句；使用完整段落。')
            if not r.get('coordinates'):notes.append('缺标记坐标；页序使用段落位置。')
            notes.append('名称匹配是本地清单匹配；未匹配不排除内容引用。具体报告名称不能确定时填0，原始题名与文献保留在JSON。')
            notes.extend(match.get('notes',[]))
            if r.get('context_audit',{}).get('status')=='layout_repaired':
                notes.append('上下文已按PDF版面与结构证据自动校正；原文和修复证据保存在JSON.context_audit。')
            vis=visual_context(d,{**r,**occ.get(r['location_id'],{})})
            if not row.get('paragraph_text') and r.get('pdf_context'):
                row['paragraph_text']=r['pdf_context'];notes.append('段落未对齐；此上下文来自PDF标记邻近文本，非已确认完整段落。')
            row.update(report_id=rid,part=classify(' '.join(chain)),section_title=chain[0] if chain else r.get('section_title'),subsection_title=' / '.join(chain[1:]) or None,
                       pdf_pages=pages[0] if len(pages)==1 else ','.join(map(str,pages)) or None,
                       confirmation_status='待核实',context_status='需人工复查',notes='\n'.join(notes),collector='paper-citations（自动导入）',reviewer=None,
                       processed_date=datetime.now().date(),rule_version='paper-citations 0.7.9',software=json.dumps(d['parser_result'].get('software',{}),ensure_ascii=False),
                       carrier={'p':'正文','cell':'表或表注','table':'表或表注','figDesc':'图或图注','note':'脚注或尾注','abstract':'正文'}.get(r.get('carrier'),'正文'),
                       report_match_status='待核实' if disputed else {'matched':'已匹配','unmatched':'未匹配'}.get(match['status'],'待核实'),report_title=title,
                       other_references='; '.join(dict.fromkeys(e.get('target_marker') or e['local_reference_id'] for e in occ[r['location_id']]['links'] if e['reference_id']!=r['report_candidate_id'])),
                       visual_context=json.dumps(vis,ensure_ascii=False) if vis else None)
            if vis:
                row['notes']+='\n图表上下文='+row['visual_context']
                if not row.get('paragraph_text'):
                    row['paragraph_text']='\n'.join(t for item in vis for k in ['captions','cells','footnotes'] for t in item[k])
                    row['notes']+='\n所在内容为图表文本串行化，不是正文段落。'
            citations.append(row)
            if rid:
                report=reports.setdefault(rid,dict(report_id=rid,target_name=target,title=title,year=r.get('year_label_from_raw'),url=match.get('matched_url'),verification_status='待核实',notes='本地清单资源候选，非正式报告编号；原文来源：'))
                report['notes']+='\n'+pid+': '+r['raw_reference']
                if report['year']!=r.get('year_label_from_raw'):report['year']=None
    if len(targets)>1:raise ValueError('Packets contain different targets; export separate workbooks')
    return dict(citations=citations,papers=papers,reports=list(reports.values()))

def export_packets(paths,template,mapping,output,overwrite=False,*,id_state=None,progress=None):
    check(id_state)
    paths=list(paths)
    output=Path(output).resolve()
    output.parent.mkdir(parents=True,exist_ok=True)
    progress=progress or Progress('填写模板',len(paths),output.with_suffix('.progress.json'))
    try:
        result=_export_packets(paths,template,mapping,output,overwrite,id_state,progress)
        progress.finish()
        return result
    except Exception:
        progress.emit('failed')
        raise

def _export_packets(paths,template,mapping,output,overwrite,id_state,progress):
    from openpyxl.utils.cell import range_boundaries,get_column_letter
    output=Path(output).resolve();template=Path(template).resolve()
    if output==template:raise ValueError('Output must not overwrite the blank template')
    if output.suffix.lower()!='.xlsx':raise ValueError('Output must be .xlsx')
    if output.exists() and not overwrite:raise FileExistsError(output)
    wb,spec=load_template(template,mapping);before=digest(template);data=collect(paths,id_state)
    def write_rows(cfg, rows, start):
        sheet=wb[cfg['sheet']]
        style_row=cfg.get('start_row',cfg.get('header_row',1)+1)
        for i,row in enumerate(rows,start):
            for col,field in cfg['_positions'].items():
                cell=sheet.cell(i,col)
                if cell.value is not None:raise ValueError(f'Template data cell is not empty: {sheet.title}!{cell.coordinate}')
                if i>style_row and not cell.has_style:cell._style=copy(sheet.cell(style_row,col)._style)
                value=row.get(field)
                if isinstance(value,str) and len(value)>32767:raise ValueError(f'Excel text limit exceeded: {field}; JSON retained, no silent truncation')
                cell.value=value
                if isinstance(value,str):cell.data_type='s'
                if field=='processed_date':cell.number_format='yyyy-mm-dd'
    offsets={cfg['sheet']:cfg.get('start_row',cfg.get('header_row',1)+1) for cfg in spec['sheets']}
    for paper in data['papers']:
        for cfg in spec['sheets']:
            if cfg['entity']=='reports':continue
            rows=[r for r in data[cfg['entity']] if r['paper_id']==paper['paper_id']]
            write_rows(cfg,rows,offsets[cfg['sheet']])
            offsets[cfg['sheet']]+=len(rows)
        # Leave the final paper pending until reports and the workbook are saved.
        if paper is not data['papers'][-1]:progress.advance(paper['paper_id'])
    for cfg in spec['sheets']:
        sheet=wb[cfg['sheet']];rows=data[cfg['entity']];start=cfg.get('start_row',cfg.get('header_row',1)+1)
        if cfg['entity']=='reports':write_rows(cfg,rows,start)
        end=start+len(rows)-1
        for table in sheet.tables.values():
            a,b,c,e=range_boundaries(table.ref)
            if b==cfg.get('header_row',1) and end>e:
                table.ref=f'{get_column_letter(a)}{b}:{get_column_letter(c)}{end}'
                if table.autoFilter:table.autoFilter.ref=table.ref
        for validation in sheet.data_validations.dataValidation:
            for rng in list(validation.ranges):
                if rng.min_row<=start<=rng.max_row and end>rng.max_row:
                    validation.add(f'{get_column_letter(rng.min_col)}{rng.max_row+1}:{get_column_letter(rng.max_col)}{end}')
    output.parent.mkdir(parents=True,exist_ok=True)
    with tempfile.NamedTemporaryFile(suffix='.xlsx',dir=output.parent,delete=False) as stream:tmp=Path(stream.name)
    try:wb.save(tmp);tmp.replace(output)
    finally:tmp.unlink(missing_ok=True)
    assert digest(template)==before,'Template unexpectedly changed'
    summary={'numbering_state':str(Path(id_state).resolve()), 'paper_ids':[p['paper_id'] for p in data['papers']],
             'record_ids':[r['record_id'] for r in data['citations']],
             'numbering':[{'source_json':p['source_json'],'source_paper_id':p['source_paper_id'],'paper_id':p['paper_id'],
                           'records':{r['source_record_id']:r['record_id'] for r in data['citations'] if r['paper_id']==p['paper_id']}} for p in data['papers']], 'output':str(output),'counts':{k:len(v) for k,v in data.items()},'template_sha256':before,'mapping_sha256':digest(mapping),
             'sources':{str(p):digest(p) for p in paths},'semantics':'candidate export, not confirmed citations'}
    output.with_suffix('.provenance.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2),encoding='utf-8')
    from .quality import write_receipt
    write_receipt(output,template,mapping,spec,data)
    wb.close()
    if data['papers']:progress.advance(data['papers'][-1]['paper_id'])
    return summary
