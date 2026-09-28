"""Native desktop interface; parsing runs in an isolated background process."""
import json
import os
from pathlib import Path
import queue
import signal
import subprocess
import sys
import threading
from datetime import datetime


def build_command(values, project):
    """Use the same public CLI and validations as the distributable tool."""
    if values.get('config'):
        command=[sys.executable,'-m','paper_citation_pipeline','run','--config',values['config'],'--resume']
        if values.get('overwrite'):command.append('--overwrite-excel')
        return command
    for name in ('input','template','target'):
        if not values.get(name,'').strip():raise ValueError('请填写 PDF 文件夹、Excel 模板和目标作者／机构。')
    cmd=[sys.executable,'-m','paper_citation_pipeline','setup',str(project),'--run','--non-interactive']
    for name in ('input','template','target','aliases','mapping','id-state','kind','grobid-url'):
        if values.get(name):cmd.extend(['--'+name,values[name]])
    if not values.get('id-state'):
        for name in ('paper-current','record-current'):
            try:number=int(values.get(name,''))
            except ValueError:raise ValueError('请填写最后已用的论文号和记录号；全新项目填 0。')
            if number<0:raise ValueError('编号不能小于 0。')
            cmd.extend(['--'+name,str(number)])
        for name in ('paper-prefix','record-prefix'):
            cmd.extend(['--'+name,values.get(name) or ('P' if name=='paper-prefix' else 'R')])
    return cmd


def progress_files(project):
    config=project/'run.json'
    if not config.exists():return []
    d=json.loads(config.read_text(encoding='utf-8'))
    out=(project/d['output']).resolve()
    xlsx=(project/d['excel']).resolve() if d.get('excel') else out/'citations.xlsx'
    return [out/'conversion-progress.json',xlsx.with_suffix('.progress.json'),xlsx.with_suffix('.quality-progress.json')]


def launch_path(path):
    if sys.platform=='darwin':subprocess.Popen(['open',str(path)])
    elif os.name=='nt':os.startfile(str(path))
    else:subprocess.Popen(['xdg-open',str(path)])


class App:
    def __init__(self, root):
        import tkinter as tk
        from tkinter import ttk, filedialog, messagebox
        self.tk,self.ttk,self.filedialog,self.messagebox=tk,ttk,filedialog,messagebox
        self.root=root;self.events=queue.Queue();self.process=None;self.running=False;self.cancelled=False
        self.project=None;self.current_config=None
        root.title('论文引用解析');root.geometry('940x820');root.minsize(820,660)
        style=ttk.Style(root);style.configure('Title.TLabel',font=('',23,'bold'));style.configure('Sub.TLabel',foreground='#586475')
        outer=ttk.Frame(root,padding=22);outer.pack(fill='both',expand=True)
        ttk.Label(outer,text='论文 → 引用证据 → Excel',style='Title.TLabel').pack(anchor='w')
        ttk.Label(outer,text='仅处理英语论文。非英语自动排除，解析、编号和填表在后台完成。',style='Sub.TLabel').pack(anchor='w',pady=(7,16))
        self.tabs=ttk.Notebook(outer);self.tabs.pack(fill='both',expand=True)
        settings=ttk.Frame(self.tabs,padding=14);self.monitor=ttk.Frame(self.tabs,padding=14)
        self.tabs.add(settings,text='  配置任务  ');self.tabs.add(self.monitor,text='  运行进度与结果  ')
        canvas=tk.Canvas(settings,highlightthickness=0);scroll=ttk.Scrollbar(settings,orient='vertical',command=canvas.yview)
        scroll.pack(side='right',fill='y');canvas.pack(side='left',fill='both',expand=True);canvas.configure(yscrollcommand=scroll.set)
        form=ttk.Frame(canvas);window=canvas.create_window((0,0),window=form,anchor='nw')
        form.bind('<Configure>',lambda e:canvas.configure(scrollregion=canvas.bbox('all')))
        canvas.bind('<Configure>',lambda e:canvas.itemconfigure(window,width=e.width));form.columnconfigure(1,weight=1)
        self.vars={};self.inputs=[];self.row=0
        self.field(form,'PDF 文件夹','input',folder=True)
        self.field(form,'Excel 空白模板','template',extension='xlsx')
        self.field(form,'目标作者／机构','target')
        ttk.Label(form,text='目标类型').grid(row=self.row,column=0,sticky='w',pady=6)
        self.vars['kind']=tk.StringVar(value='机构');box=ttk.Combobox(form,textvariable=self.vars['kind'],values=['机构','个人作者'],state='readonly');box.grid(row=self.row,column=1,sticky='ew',padx=10);self.inputs.append(box);self.row+=1
        self.field(form,'名称变体 JSON（可选）','aliases',extension='json')
        self.field(form,'模板映射 JSON（可选）','mapping',extension='json')
        ttk.Button(form,text='在窗口中配置模板字段…',command=self.mapping_editor).grid(row=self.row,column=1,sticky='w',padx=10,pady=5);self.row+=1
        self.field(form,'项目保存位置','parent',folder=True,default=str(Path.cwd()))
        ttk.Separator(form).grid(row=self.row,column=0,columnspan=3,sticky='ew',pady=12);self.row+=1
        self.field(form,'已有编号文件（可选）','id-state',extension='json')
        ttk.Label(form,text='无编号文件时填写下方两项。数字指最后已用号码；新项目填 0。',style='Sub.TLabel',wraplength=670).grid(row=self.row,column=0,columnspan=3,sticky='w',pady=5);self.row+=1
        self.field(form,'最后已用论文号','paper-current');self.field(form,'最后已用记录号','record-current')
        self.field(form,'论文编号前缀','paper-prefix',default='P');self.field(form,'记录编号前缀','record-prefix',default='R')
        self.field(form,'GROBID 服务地址','grobid-url',default='http://localhost:8070')
        ttk.Label(form,text='标准模板自动匹配字段；其他模板可点击上方按钮配置。原模板保留，结果另存。',style='Sub.TLabel',wraplength=660).grid(row=self.row,column=0,columnspan=3,sticky='w',pady=10)
        actions=ttk.Frame(outer);actions.pack(fill='x',pady=(14,0))
        self.start_button=ttk.Button(actions,text='开始解析并生成 Excel',command=self.start);self.start_button.pack(side='right')
        self.resume_button=ttk.Button(actions,text='继续已有项目…',command=self.resume);self.resume_button.pack(side='left')
        self.check_button=ttk.Button(actions,text='检查运行环境',command=self.doctor);self.check_button.pack(side='left',padx=8)
        self.status=tk.StringVar(value='等待选择文件');ttk.Label(self.monitor,textvariable=self.status,wraplength=800).pack(anchor='w',pady=(0,18))
        self.bars=[];self.labels=[]
        for title in ('PDF → Markdown / JSON','填写 Excel 模板'):
            ttk.Label(self.monitor,text=title,font=('',14,'bold')).pack(anchor='w')
            label=tk.StringVar(value='等待开始');ttk.Label(self.monitor,textvariable=label).pack(anchor='w',pady=5)
            bar=ttk.Progressbar(self.monitor,maximum=100);bar.pack(fill='x',pady=(0,20));self.bars.append(bar);self.labels.append(label)
        buttons=ttk.Frame(self.monitor);buttons.pack(fill='x',pady=(0,12))
        self.open_excel=ttk.Button(buttons,text='打开 Excel',command=lambda:self.open_result(True),state='disabled');self.open_excel.pack(side='left')
        self.open_folder=ttk.Button(buttons,text='打开结果文件夹',command=lambda:self.open_result(False),state='disabled');self.open_folder.pack(side='left',padx=8)
        self.stop_button=ttk.Button(buttons,text='停止本次处理',command=self.stop,state='disabled');self.stop_button.pack(side='right')
        ttk.Label(self.monitor,text='运行日志（候选引用保留待确认状态）',style='Sub.TLabel').pack(anchor='w')
        from tkinter.scrolledtext import ScrolledText
        self.log=ScrolledText(self.monitor,height=12,wrap='word',font=('',11),state='disabled');self.log.pack(fill='both',expand=True,pady=6)
        root.protocol('WM_DELETE_WINDOW',self.close);root.after(250,self.poll)

    def field(self,parent,label,key,folder=False,extension=None,default=''):
        ttk,tk=self.ttk,self.tk
        ttk.Label(parent,text=label).grid(row=self.row,column=0,sticky='w',pady=6)
        variable=tk.StringVar(value=default);self.vars[key]=variable
        entry=ttk.Entry(parent,textvariable=variable);entry.grid(row=self.row,column=1,sticky='ew',padx=10,pady=6);self.inputs.append(entry)
        if folder or extension:
            def browse():
                if self.running:return
                path=self.filedialog.askdirectory(parent=self.root) if folder else self.filedialog.askopenfilename(parent=self.root,filetypes=[(extension.upper(),'*.'+extension)])
                if path:variable.set(path)
            button=ttk.Button(parent,text='选择…',command=browse);button.grid(row=self.row,column=2);self.inputs.append(button)
        self.row+=1

    def mapping_editor(self):
        if self.running:return
        path=self.vars['template'].get()
        if not path:self.messagebox.showinfo('选择模板','请先选择 Excel 模板。');return
        try:MappingEditor(self,path)
        except Exception as e:self.messagebox.showerror('无法读取模板',str(e))

    def start(self,config=None):
        if self.running:return
        try:
            values={k:v.get().strip() for k,v in self.vars.items()}
            values['kind']={'机构':'organization','个人作者':'person'}.get(values['kind'],values['kind'])
            if config:
                values={'config':str(config)}
                d=json.loads(config.read_text(encoding='utf-8'));self.project=config.parent
                xlsx=(self.project/d['excel']).resolve() if d.get('excel') else (self.project/d['output']/'citations.xlsx').resolve()
                if xlsx.exists():
                    if not self.messagebox.askyesno('已有结果','此项目已有 Excel。是否重跑并覆盖这个结果文件？空白模板不会更改。'):return
                    values['overwrite']=True
            else:
                parent=Path(values['parent']).expanduser()
                if not parent.is_dir():raise ValueError('请选择有效的项目保存文件夹。')
                self.project=parent/('论文解析-'+datetime.now().strftime('%Y%m%d-%H%M%S-%f'))
            cmd=build_command(values,self.project)
        except Exception as e:self.messagebox.showerror('配置未完成',str(e));return
        self.running=True;self.cancelled=False;self.current_config=config
        self.tabs.select(self.monitor);self.status.set('正在准备任务…')
        for bar,label in zip(self.bars,self.labels):bar['value']=0;label.set('等待开始')
        self.log.configure(state='normal');self.log.delete('1.0','end');self.log.configure(state='disabled')
        for button in (self.start_button,self.resume_button,self.check_button,self.open_excel,self.open_folder):button.configure(state='disabled')
        self.stop_button.configure(state='normal')
        threading.Thread(target=self.worker,args=(cmd,),daemon=True).start()

    def worker(self,cmd):
        try:
            kwargs={'stdout':subprocess.PIPE,'stderr':subprocess.STDOUT,'stdin':subprocess.DEVNULL,'text':True,'encoding':'utf-8','errors':'replace','env':dict(os.environ,PYTHONUNBUFFERED='1')}
            if os.name=='nt':kwargs['creationflags']=subprocess.CREATE_NO_WINDOW|subprocess.CREATE_NEW_PROCESS_GROUP
            else:kwargs['start_new_session']=True
            if getattr(self,'api_key',''):kwargs['env']['PAPER_CITATIONS_API_KEY']=self.api_key
            self.process=subprocess.Popen(cmd,**kwargs)
            if self.cancelled:self.terminate()
            for line in self.process.stdout:self.events.put(('log',line))
            self.events.put(('done',self.process.wait()))
        except Exception as e:self.events.put(('log',str(e)+'\n'));self.events.put(('done',2))

    def resume(self):
        path=self.filedialog.askopenfilename(title='选择项目 run.json',filetypes=[('运行配置','*.json')])
        if path:self.start(Path(path).resolve())

    def doctor(self):
        self.check_button.configure(state='disabled')
        url=self.vars['grobid-url'].get()
        def work():
            try:
                result=subprocess.run([sys.executable,'-m','paper_citation_pipeline','doctor','--grobid-url',url],capture_output=True,text=True,timeout=30,**({'creationflags':subprocess.CREATE_NO_WINDOW} if os.name=='nt' else {}))
                self.events.put(('doctor',(result.returncode,result.stdout or result.stderr)))
            except Exception as e:self.events.put(('doctor',(2,str(e))))
        threading.Thread(target=work,daemon=True).start()

    def poll(self):
        try:
            while True:
                event,value=self.events.get_nowait()
                if event=='log':
                    self.log.configure(state='normal');self.log.insert('end',value);self.log.see('end');self.log.configure(state='disabled')
                elif event=='doctor':
                    self.check_button.configure(state='normal')
                    (self.messagebox.showinfo if value[0]==0 else self.messagebox.showerror)('环境已就绪' if value[0]==0 else '需要检查运行环境',value[1])
                elif event=='done':
                    self.running=False;self.process=None
                    self.start_button.configure(state='normal');self.resume_button.configure(state='normal');self.check_button.configure(state='normal');self.stop_button.configure(state='disabled')
                    self.status.set('已停止。可使用“继续已有项目”复用已完成解析和编号。' if self.cancelled else ('完成：Excel 与证据文件已保存。' if value==0 else '任务未完整完成，请查看下方日志；已有证据保留。'))
                    self.refresh_result_buttons()
        except queue.Empty:pass
        if self.project:
            try:
                for i,path in enumerate(progress_files(self.project)):
                    if not path.exists():continue
                    d=json.loads(path.read_text(encoding='utf-8'));total=d['total'];done=d['processed']
                    self.bars[i]['value']=100*done/total if total else 0
                    self.labels[i].set(f'{done} / {total} 篇 · 成功 {d["succeeded"]} · 排除 {d.get("excluded",0)} · 失败 {d["failed"]}'+(' · '+d['current_paper'] if d.get('current_paper') else ''))
                    if self.running:self.status.set('正在解析与生成结果，请保持窗口打开。')
            except (OSError,ValueError,KeyError):pass
        self.root.after(300,self.poll)

    def result_paths(self):
        if not self.project or not (self.project/'run.json').exists():return None,None
        cfg=json.loads((self.project/'run.json').read_text(encoding='utf-8'));out=(self.project/cfg['output']).resolve()
        return out,(self.project/cfg['excel']).resolve() if cfg.get('excel') else out/'citations.xlsx'

    def refresh_result_buttons(self):
        try:
            out,xlsx=self.result_paths()
            if out and out.exists():self.open_folder.configure(state='normal')
            if xlsx and xlsx.exists():self.open_excel.configure(state='normal')
        except (OSError,ValueError,KeyError):pass

    def open_result(self,excel):
        out,xlsx=self.result_paths();path=xlsx if excel else out
        if path and path.exists():launch_path(path)

    def terminate(self):
        p=self.process
        if p is None or p.poll() is not None:return
        try:
            if os.name=='nt':subprocess.run(['taskkill','/PID',str(p.pid),'/T','/F'],capture_output=True,creationflags=subprocess.CREATE_NO_WINDOW)
            else:os.killpg(p.pid,signal.SIGTERM)
        except ProcessLookupError:pass

    def stop(self):
        if not self.running:return
        self.cancelled=True;self.terminate();self.status.set('正在停止…已分配编号会保留。')

    def close(self):
        if self.running:
            if not self.messagebox.askyesno('任务正在运行','关闭窗口会停止本次任务，是否继续？'):return
            self.stop()
        self.root.destroy()


class MappingEditor:
    def __init__(self,app,path):
        from openpyxl import load_workbook
        from .config import read_json
        from .excel import FIELDS
        self.app=app;self.path=Path(path);self.fields=FIELDS
        self.wb=load_workbook(path,read_only=True)
        self.known={s['sheet']:s for s in read_json(Path(__file__).parent/'examples/unep-template-map.json')['sheets']}
        self.spec={'sheets':[]};self.index=0
        tk,ttk=app.tk,app.ttk
        self.top=tk.Toplevel(app.root);self.top.title('模板字段配置');self.top.geometry('780x650')
        self.title=tk.StringVar();ttk.Label(self.top,textvariable=self.title,font=('',16,'bold')).pack(anchor='w',padx=15,pady=12)
        row=ttk.Frame(self.top);row.pack(fill='x',padx=15)
        self.entity=tk.StringVar();ttk.Label(row,text='工作表内容').pack(side='left')
        ttk.Combobox(row,textvariable=self.entity,values=['citations','papers','reports','skip'],state='readonly',width=15).pack(side='left',padx=8)
        self.header=tk.StringVar(value='1');self.start=tk.StringVar(value='2')
        ttk.Label(row,text='表头行').pack(side='left');ttk.Entry(row,textvariable=self.header,width=4).pack(side='left')
        ttk.Label(row,text='首条数据行').pack(side='left',padx=8);ttk.Entry(row,textvariable=self.start,width=4).pack(side='left')
        ttk.Button(row,text='加载列',command=self.columns).pack(side='right')
        ttk.Label(self.top,text='每列选择对应字段；skip 表示明确保留空白。更换类型或行号后点击“加载列”。').pack(anchor='w',padx=15,pady=10)
        canvas=tk.Canvas(self.top,highlightthickness=0);scroll=ttk.Scrollbar(self.top,command=canvas.yview);scroll.pack(side='right',fill='y');canvas.pack(fill='both',expand=True,padx=15)
        self.body=ttk.Frame(canvas);win=canvas.create_window(0,0,window=self.body,anchor='nw');canvas.configure(yscrollcommand=scroll.set)
        self.body.bind('<Configure>',lambda e:canvas.configure(scrollregion=canvas.bbox('all')));canvas.bind('<Configure>',lambda e:canvas.itemconfigure(win,width=e.width))
        ttk.Button(self.top,text='保存本表并继续',command=self.next).pack(pady=14)
        self.top.protocol('WM_DELETE_WINDOW',self.close);self.show()

    def show(self):
        ws=self.wb.worksheets[self.index];self.title.set(f'{self.index+1}/{len(self.wb.worksheets)} · {ws.title}')
        self.entity.set(self.known.get(ws.title,{}).get('entity','skip' if '填写' in ws.title else 'citations'))
        self.header.set('1');self.start.set('2');self.columns()

    def columns(self):
        for child in self.body.winfo_children():child.destroy()
        self.options=[];self.loaded=None
        entity=self.entity.get()
        if entity=='skip':return
        try:
            header=int(self.header.get());start=int(self.start.get())
            if header<1 or start<=header:raise ValueError('表头行必须为正数，数据起始行必须在表头之后。')
            ws=self.wb.worksheets[self.index];known=self.known.get(ws.title,{})
            lookup=known.get('columns',{}) if known.get('entity')==entity else {}
            headers=[str(c.value) for c in ws[header] if c.value is not None]
            if len(headers)!=len(set(headers)):raise ValueError('表头有重名，请先修改模板。')
            for i,title in enumerate(headers):
                variable=self.app.tk.StringVar(value=lookup.get(title,title if title in self.fields[entity] else '请选择'))
                self.app.ttk.Label(self.body,text=title,wraplength=350).grid(row=i,column=0,sticky='w',pady=5)
                self.app.ttk.Combobox(self.body,textvariable=variable,values=['skip']+sorted(self.fields[entity]),state='readonly',width=32).grid(row=i,column=1,padx=15,pady=5)
                self.options.append((title,variable))
            self.loaded=(entity,header,start)
        except Exception as e:self.app.messagebox.showerror('无法加载列',str(e),parent=self.top)

    def next(self):
        if self.index>=len(self.wb.worksheets):
            self.save();return
        entity=self.entity.get()
        if entity!='skip':
            try:
                if self.loaded!=(entity,int(self.header.get()),int(self.start.get())):raise ValueError('请先点击“加载列”。')
                if any(v.get()=='请选择' for _,v in self.options):raise ValueError('请为所有列选择字段或明确选择 skip。')
                columns={title:v.get() for title,v in self.options if v.get()!='skip'}
                if not columns:raise ValueError('至少选择一列，或将工作表内容设为 skip。')
                self.spec['sheets'].append(dict(sheet=self.wb.worksheets[self.index].title,entity=entity,header_row=int(self.header.get()),start_row=int(self.start.get()),columns=columns))
            except Exception as e:self.app.messagebox.showerror('映射未完成',str(e),parent=self.top);return
        self.index+=1
        if self.index<len(self.wb.worksheets):self.show();return
        if not self.spec['sheets']:
            self.app.messagebox.showerror('未选择字段','至少需要配置一张工作表。',parent=self.top);self.index=0;self.show();return
        self.save()

    def save(self):
        dest=self.app.filedialog.asksaveasfilename(parent=self.top,title='保存模板字段映射',defaultextension='.json',initialfile='template-map.json',filetypes=[('JSON','*.json')])
        if not dest:return
        try:
            Path(dest).write_text(json.dumps(self.spec,ensure_ascii=False,indent=2),encoding='utf-8')
            self.app.vars['mapping'].set(dest);self.close()
        except OSError as e:self.app.messagebox.showerror('保存失败',str(e),parent=self.top)

    def close(self):self.wb.close();self.top.destroy()


def main():
    import tkinter as tk
    from .workbench import Workbench
    root=tk.Tk();Workbench(root);root.mainloop()

if __name__=='__main__':main()
