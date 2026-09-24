"""Audit completed workbooks against export receipts, without reading paper content."""
from datetime import date, datetime
import hashlib
import html
import json
from pathlib import Path
from openpyxl import load_workbook
from .config import digest, read_json
from .numbering import atomic_json
from .progress import Progress

def signature(value):
    if value is None or value == '': value = None
    elif isinstance(value, (datetime,date)): value = value.isoformat()[:10]
    elif isinstance(value,float) and value.is_integer(): value=int(value)
    return {'type':type(value).__name__, 'length':len(value) if isinstance(value,str) else None,
            'sha256':hashlib.sha256(json.dumps(value,ensure_ascii=False,sort_keys=True).encode()).hexdigest(),
            'blank':value is None}


def write_receipt(output,template,mapping,spec,data):
    """Expected values come from logical export data, never from the filled workbook."""
    blank=load_workbook(template)
    sheets=[]
    for cfg in spec['sheets']:
        ws=blank[cfg['sheet']];start=cfg.get('start_row',cfg.get('header_row',1)+1)
        rows=[]
        for index,row in enumerate(data[cfg['entity']],start):
            cells=[]
            for col,field in cfg['_positions'].items():
                cells.append({'column':col,'field':field,'signature':signature(row.get(field))})
            rows.append({'row':index,'paper_id':row.get('paper_id'),'cells':cells})
        sheets.append({'sheet':cfg['sheet'],'entity':cfg['entity'],'header_row':cfg.get('header_row',1),
                       'start_row':start,'rows':rows,'headers':{str(c.column):c.value for c in ws[cfg.get('header_row',1)] if c.value is not None},
                       'columns':cfg['columns']})
    receipt={'version':1,'template_sha256':digest(template),'mapping_sha256':digest(mapping),
             'output_sha256':digest(output),'sheets':sheets,'paper_ids':[p['paper_id'] for p in data['papers']],
             'counts':{k:len(v) for k,v in data.items()}}
    blank.close();atomic_json(Path(output).with_suffix('.fill-manifest.json'),receipt)


def inspect_workbook(workbook,expected_papers=None):
    workbook=Path(workbook);receipt=read_json(workbook.with_suffix('.fill-manifest.json'))
    wb=load_workbook(workbook,data_only=False);checks=[]
    def add(status,kind,**details):
        checks.append({'id':f'q{len(checks)+1:05d}','status':status,'kind':kind,**details})
    n=len(receipt['paper_ids'])
    add('pass' if expected_papers is None or expected_papers==n else 'error','paper_count',expected=expected_papers if expected_papers is not None else n,actual=n)
    add('pass' if digest(workbook)==receipt['output_sha256'] else 'warning','workbook_changed_since_export')
    canonical=read_json(Path(__file__).parent/'examples/unep-template-map.json')
    known={x['entity']:x['columns'] for x in canonical['sheets']}
    identifiers={};foreign=[]
    for cfg in receipt['sheets']:
        name=cfg['sheet'];entity=cfg['entity']
        if name not in wb.sheetnames:add('error','missing_sheet',sheet=name);continue
        ws=wb[name];positions={}
        for col,header in cfg['headers'].items():
            actual=ws.cell(cfg['header_row'],int(col)).value
            add('pass' if actual==header else 'error','header',sheet=name,column=int(col),header=str(header))
            field=cfg['columns'].get(header)
            if field:
                positions[field]=int(col)
                standard=known.get(entity,{}).get(header)
                add('pass' if standard is None or field==standard else 'error','field_mapping',sheet=name,header=header,field=field,expected_field=standard)
            elif str(header).endswith('*'):add('warning','required_header_unmapped',sheet=name,header=str(header))
        start=cfg['start_row'];end=start+len(cfg['rows'])
        actual_rows=sum(any(ws.cell(r,c).value is not None for c in positions.values()) for r in range(start,ws.max_row+1))
        add('pass' if actual_rows==len(cfg['rows']) else 'error','row_count',sheet=name,expected=len(cfg['rows']),actual=actual_rows)
        required={field for h,field in cfg['columns'].items() if h.endswith('*')}
        idfield={'papers':'paper_id','citations':'record_id','reports':'report_id'}[entity];ids=[]
        for row in cfg['rows']:
            matched=0;paper=row.get('paper_id')
            for expected in row['cells']:
                cell=ws.cell(row['row'],expected['column']);sig=signature(cell.value);field=expected['field']
                where=dict(sheet=name,row=row['row'],column=expected['column'],field=field,paper_id=paper)
                if sig!=expected['signature']:add('error','cell_mismatch',expected=expected['signature'],actual=sig,**where)
                else:matched+=1
                if cell.data_type=='f':add('error','unexpected_formula',**where)
                if sig['blank'] and field in required:add('warning','required_blank_in_source' if expected['signature']['blank'] else 'required_value_lost',**where)
                for dv in ws.data_validations.dataValidation:
                    if dv.type=='list' and cell.coordinate in dv and dv.formula1 and dv.formula1.startswith('"') and not sig['blank']:
                        if str(cell.value) not in dv.formula1.strip('"').split(','):add('error','invalid_enum',**where)
            add('pass','cells_compared',sheet=name,row=row['row'],paper_id=paper,matched=matched,total=len(row['cells']))
        # Include unexpected rows in duplicate/foreign-key checks.
        for r in range(start,ws.max_row+1):
            if idfield in positions:
                value=ws.cell(r,positions[idfield]).value
                if value is not None:ids.append(str(value))
            if entity=='citations':
                for field in ('paper_id','report_id'):
                    if field in positions:
                        value=ws.cell(r,positions[field]).value
                        if value is not None:foreign.append((name,r,field,str(value)))
        if idfield in positions:
            add('pass' if len(ids)==len(set(ids)) else 'error','duplicate_identifiers',sheet=name,field=idfield,count=len(ids),unique=len(set(ids)))
            identifiers[idfield]=set(ids)
    for sheet,row,field,value in foreign:
        # Paper list is known even when a custom template only exports citations.
        valid=identifiers.get(field,set(receipt['paper_ids']) if field=='paper_id' else None)
        if valid is not None and value not in valid:add('error','broken_foreign_key',sheet=sheet,row=row,field=field)
    wb.close();return receipt,checks


def review_workbook(workbook,cfg=None,expected_papers=None):
    # cfg is accepted only for old integrations; it cannot enable model calls.
    workbook=Path(workbook);receipt,checks=inspect_workbook(workbook,expected_papers)
    report={'version':2,'scope':'local structural checks only','counts':receipt['counts'],
            'checks':checks,'ai_enabled':False,'ai_checks':[],'status':'running'}
    progress=Progress('保存结果',len(receipt['paper_ids']),workbook.with_suffix('.quality-progress.json'))
    for pid in receipt['paper_ids']:
        progress.advance(pid,failed=any(c['status']=='error' and c.get('paper_id')==pid for c in checks))
    statuses=[c['status'] for c in checks]
    report['status']='error' if 'error' in statuses else 'needs_attention' if 'warning' in statuses else 'passed'
    progress.finish()
    atomic_json(workbook.with_suffix('.quality.json'),report)
    write_html(workbook.with_suffix('.quality.html'),report)
    return 1 if report['status']=='error' else 0


def write_html(path,report):
    labels={'paper_count':'论文数量','workbook_changed_since_export':'导出后文件有变化','missing_sheet':'缺少工作表',
            'header':'表头变化','field_mapping':'字段映射','required_header_unmapped':'必填列未映射',
            'row_count':'记录数量','cell_mismatch':'单元格与预期不符','unexpected_formula':'意外公式',
            'required_blank_in_source':'源数据缺少必填内容','required_value_lost':'必填内容导出丢失',
            'invalid_enum':'内容不符合下拉选项','duplicate_identifiers':'重复编号','broken_foreign_key':'编号关联异常'}
    states={'passed':'填表质检通过','needs_attention':'有待处理项','error':'发现输出异常','failed':'程序检查未完成'}
    issues=[c for c in report['checks'] if c['status']!='pass']
    rows=[]
    for c in issues:
        where=str(c.get('sheet','整个任务'))
        if c.get('row'):where+=' · 第'+str(c['row'])+'行'
        if c.get('column'):where+=' · 第'+str(c['column'])+'列'
        detail=c.get('field',c.get('header',''))
        if c['kind'] in {'row_count','paper_count'}:detail=f"预期 {c['expected']}，实际 {c['actual']}"
        rows.append('<tr><td>'+html.escape(labels.get(c['kind'],c['kind']))+'</td><td>'+html.escape(where)+'</td><td>'+html.escape(str(detail))+'</td></tr>')
    content='<!doctype html><meta charset="utf-8"><title>填表结果质检</title><style>body{font:15px system-ui;color:#183142;background:#f3f6f8;max-width:1080px;margin:36px auto;padding:24px}table{border-collapse:collapse;width:100%;background:white}th,td{text-align:left;border-bottom:1px solid #dde5e9;padding:12px}pre{white-space:pre-wrap;overflow-wrap:anywhere}h1{color:#087e75}</style>'
    content+='<h1>'+states.get(report['status'],report['status'])+'</h1><p>论文 '+str(report['counts']['papers'])+' 篇 · 引用记录 '+str(report['counts']['citations'])+' 条 · 程序待处理项 '+str(len(issues))+' 项</p>'
    content+='<p>仅检查填表结果与映射，不检查原文引用。源数据必填项缺失与导出丢失分别标记；可选空值不报错。</p>'
    if report.get('error'):content+='<p>'+html.escape(report['error'])+'；Excel已保留。</p>'
    content+='<table><thead><tr><th>问题</th><th>位置</th><th>字段或数量</th></tr></thead><tbody>'+(''.join(rows) or '<tr><td colspan="3">程序检查未发现异常</td></tr>')+'</tbody></table>'
    content+='<details><summary>完整检查数据</summary><pre>'+html.escape(json.dumps(report,ensure_ascii=False,indent=2))+'</pre></details>'
    Path(path).write_text(content,encoding='utf-8')
