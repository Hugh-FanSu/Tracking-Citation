"""First-run project wizard; existing PDFs, templates, and ledgers stay in place."""
import json
import os
from pathlib import Path
import shutil
import sys
import tempfile
from .config import read_json, target_config
from .excel import FIELDS, load_template
from .numbering import check, initialize

HERE = Path(__file__).resolve().parent


def ask(value, label, interactive, default=None, optional=False):
    if value is not None:
        return value
    if not interactive:
        if optional: return default
        raise ValueError(f'Missing {label}; supply the corresponding setup option')
    suffix = f' [{default}]' if default is not None else ''
    answer = input(label + suffix + ': ').strip()
    if answer: return answer
    if default is not None or optional: return default
    raise ValueError(f'{label} is required')


def template_mapping(template, interactive):
    """Match exact known headers, or obtain explicit user choices. No fuzzy inference."""
    from openpyxl import load_workbook
    if template.suffix.lower() != '.xlsx': raise ValueError('Template must be a local .xlsx file')
    wb = load_workbook(template)
    known = {s['sheet']: s for s in read_json(HERE/'examples/unep-template-map.json')['sheets']}
    result = {'sheets': []}
    try:
        for ws in wb.worksheets:
            if ws.title in {'填写说明','填写示例（虚构）','填写示例虚构'}: continue
            default = known.get(ws.title, {}).get('entity')
            if ws.title in FIELDS: default = ws.title
            if not default:
                if not interactive:
                    raise ValueError(f'Unknown sheet {ws.title!r}; use interactive setup or --mapping with an explicit map')
                print(f'\n工作表：{ws.title}。选择 citations / papers / reports；输入 skip 保留且不填。')
                entity = ask(None,'工作表类型',True)
                if entity == 'skip': continue
            else: entity = default
            if entity not in FIELDS: raise ValueError('Unknown entity: ' + entity)
            header = int(ask(None,f'{ws.title} 表头行号',interactive,1,optional=True))
            start = int(ask(None,f'{ws.title} 首条数据行号',interactive,header+1,optional=True))
            if header < 1 or start <= header: raise ValueError('Invalid header or start row')
            lookup = known.get(ws.title, {}).get('columns', {}) if entity == default else {}
            cfg = dict(sheet=ws.title, entity=entity, header_row=header, start_row=start, columns={})
            headers = [str(c.value) for c in ws[header] if c.value is not None]
            if len(headers) != len(set(headers)): raise ValueError(f'Duplicate headers in {ws.title}')
            for title in headers:
                field = lookup.get(title) or (title if title in FIELDS[entity] else None)
                if field is None:
                    if not interactive:
                        raise ValueError(f'Unknown column {ws.title}/{title}; use interactive setup or --mapping')
                    print('可用字段：' + ', '.join(sorted(FIELDS[entity])))
                    field = ask(None,f'列「{title}」对应字段（skip=保持空白）',True)
                    if field == 'skip': continue
                if field not in FIELDS[entity]: raise ValueError(f'Unknown field {entity}.{field}')
                cfg['columns'][title] = field
            if cfg['columns']: result['sheets'].append(cfg)
        if not result['sheets']: raise ValueError('No template columns mapped')
        return result
    finally:
        wb.close()


def prepare(args):
    interactive = not args.non_interactive and sys.stdin.isatty()
    root = args.directory.expanduser().resolve()
    # Stage in a sibling directory and publish only a complete project.
    if root.exists(): raise FileExistsError('Project already exists; use its run.json or choose a new directory: ' + str(root))
    source = Path(ask(args.input,'PDF文件夹 (--input)',interactive)).expanduser().resolve()
    template = Path(ask(args.template,'本地Excel空白模板 (--template)',interactive)).expanduser().resolve()
    name = ask(args.target,'目标被引作者或机构 (--target)',interactive)
    kind = ask(args.kind,'目标类型 person / organization (--kind)',interactive,'organization',True)
    alias = ask(args.aliases,'本地名称变体JSON，留空仅用目标名称 (--aliases)',interactive,optional=True)
    alias = Path(alias).expanduser().resolve() if alias else None
    target_config(name,alias,kind)
    if not source.is_dir() or not any(p.is_file() and p.suffix.lower()=='.pdf' for p in source.iterdir()):
        raise ValueError('PDF folder does not exist or contains no PDFs')
    if not template.is_file(): raise ValueError('Template not found: ' + str(template))
    if args.id_state and (args.paper_current is not None or args.record_current is not None):
        raise ValueError('Use an existing --id-state OR initial counters, not both')
    state = Path(args.id_state).expanduser().resolve() if args.id_state else None
    if state: check(state)
    else:
        print('编号为最后已用号码；只有新项目才填0。已有编号文件请改用 --id-state。',flush=True)
        paper = int(ask(args.paper_current,'最后已用论文号 (--paper-current)',interactive))
        record = int(ask(args.record_current,'最后已用记录号 (--record-current)',interactive))
    spec = read_json(args.mapping) if args.mapping else template_mapping(template,interactive)
    root.parent.mkdir(parents=True,exist_ok=True)
    temp = Path(tempfile.mkdtemp(prefix='.paper-citations-setup-',dir=root.parent))
    try:
        assets = temp/'assets';assets.mkdir()
        shutil.copy2(template,assets/'template.xlsx')
        (assets/'mapping.json').write_text(json.dumps(spec,ensure_ascii=False,indent=2),encoding='utf-8')
        wb,_ = load_template(assets/'template.xlsx',assets/'mapping.json');wb.close()
        if alias: shutil.copy2(alias,assets/'aliases.json')
        if not state:
            initialize(temp/'numbering.json',paper,record,args.paper_prefix,args.record_prefix,args.width)
        config = dict(input=os.path.relpath(source,root),output='output',target=name,kind=kind,
                      template='assets/template.xlsx',mapping='assets/mapping.json',
                      id_state=os.path.relpath(state,root) if state else 'numbering.json',
                      grobid_url=args.grobid_url,resume=True)
        if alias:config['aliases']='assets/aliases.json'
        (temp/'run.json').write_text(json.dumps(config,ensure_ascii=False,indent=2),encoding='utf-8')
        shutil.copy2(HERE/'assets/compose.yaml',temp/'compose.yaml')
        (temp/'START.md').write_text('运行 `paper-citations run --config run.json`。需要本地GROBID时运行 `docker compose up -d`；运行前可用 `paper-citations doctor` 检查。\n'
                                    'PDF目录由run.json中的input指定。原模板和PDF不改写，结果在output/citations.xlsx，证据在output/json，阅读稿在output/md。\n'
                                    '新批次可改input和output，并复用id_state。已有输出需明确 --overwrite-excel；中断后可用 --resume。\n'
                                    'numbering.json及分配关系需随项目备份，不能重置；候选记录不等于已确认引用。\n',encoding='utf-8')
        if root.exists():raise FileExistsError(root)
        temp.rename(root)
    finally:
        if temp.exists():shutil.rmtree(temp)
    print('项目已就绪：' + str(root),flush=True)
    print('运行：paper-citations run --config ' + json.dumps(str(root/'run.json'),ensure_ascii=False),flush=True)
    return root/'run.json'
